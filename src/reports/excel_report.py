"""Excel report — KPIs, vendor scorecard, aging."""
from __future__ import annotations
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def _style_header(ws, row: int = 1) -> None:
    fill = PatternFill("solid", start_color="0FAAFF")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def build_excel(path: Path, *, kpis: dict, scorecard: pd.DataFrame, aging: pd.DataFrame) -> Path:
    wb = Workbook()

    # KPI sheet
    ws = wb.active
    ws.title = "KPIs"
    ws.append(["Metric", "Value"])
    for k, v in kpis.items():
        ws.append([k.replace("_", " ").title(), v])
    _style_header(ws)

    # Vendor scorecard
    ws2 = wb.create_sheet("Vendor Scorecard")
    if not scorecard.empty:
        ws2.append(list(scorecard.columns))
        for row in scorecard.itertuples(index=False):
            ws2.append(list(row))
        _style_header(ws2)

    # Aging
    ws3 = wb.create_sheet("Aging")
    if not aging.empty:
        cols = ["invoice_id", "vendor_id", "amount_eur", "days_overdue", "bucket"]
        cols = [c for c in cols if c in aging.columns]
        ws3.append(cols)
        for row in aging[cols].itertuples(index=False):
            ws3.append(list(row))
        _style_header(ws3)

    wb.save(path)
    return path
